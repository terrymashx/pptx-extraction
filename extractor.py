from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def extract_pptx(data: bytes, filename: str) -> dict[str, Any]:
    prs = Presentation(io.BytesIO(data))

    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        slides.append(_extract_slide(idx, slide))

    return {
        "filename": filename,
        "slide_count": len(slides),
        "slides": slides,
    }


def _extract_slide(index: int, slide) -> dict[str, Any]:
    title = _get_title(slide)
    text_blocks: list[str] = []
    tables: list[list[list[str]]] = []

    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            tables.append(_extract_table(shape.table))
            continue

        if not shape.has_text_frame:
            continue

        if shape == slide.shapes.title:
            continue

        block = _extract_text_frame(shape.text_frame)
        if block:
            text_blocks.append(block)

    return {
        "slide_number": index,
        "title": title,
        "text_blocks": text_blocks,
        "tables": tables,
    }


def _get_title(slide) -> str | None:
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        return slide.shapes.title.text_frame.text.strip() or None
    return None


def _extract_text_frame(tf) -> str:
    lines = []
    for para in tf.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def _extract_table(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells = [cell.text_frame.text.strip() for cell in row.cells]
        rows.append(cells)
    return rows
